import csv
import hashlib
import json
import os
from typing import Any
from uuid import uuid4

from fastapi import UploadFile
from langchain_core.documents import Document

from app.services.langchain_vector_store import get_vector_store
from app.services.chunking import chunk_text

ALLOWED_EXTENSIONS = {"pdf", "docx", "md", "txt", "csv", "json"}


def _get_extension(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lstrip(".").lower()


def _validate_extension(filename: str) -> str:
    extension = _get_extension(filename)
    if not extension:
        raise ValueError("Uploaded document must include a file extension.")
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(
            f"Unsupported document type '{extension}'. Supported types: {', '.join(sorted(ALLOWED_EXTENSIONS))}."
        )
    return extension


def _make_chunk_id(source: str, chunk_index: int, chunk_text: str) -> str:
    return hashlib.md5(f"{source}:{chunk_index}:{chunk_text}".encode("utf-8", errors="replace")).hexdigest()


async def extract_text_from_upload(file: UploadFile) -> str:
    extension = _validate_extension(file.filename or "")
    file.file.seek(0)

    if extension == "pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("PDF ingestion requires the 'pypdf' package.") from exc

        reader = PdfReader(file.file)
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)

    if extension == "docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("DOCX ingestion requires the 'python-docx' package.") from exc

        document = Document(file.file)
        return "\n\n".join(paragraph.text for paragraph in document.paragraphs)

    if extension == "csv":
        file.file.seek(0)
        content = file.file.read().decode("utf-8", errors="replace")
        rows = []
        for row in csv.reader(content.splitlines()):
            rows.append(", ".join(row))
        return "\n".join(rows)

    if extension == "json":
        file.file.seek(0)
        try:
            parsed = json.load(file.file)
        except json.JSONDecodeError as exc:
            raise ValueError("Uploaded JSON is invalid.") from exc
        return json.dumps(parsed, indent=2, ensure_ascii=False)

    # md or txt or other plain text-like files
    file.file.seek(0)
    return file.file.read().decode("utf-8", errors="replace")


async def ingest_upload(
    file: UploadFile,
    source: str | None = None,
    metadata: dict[str, Any] | None = None,
    collection_name: str = "docs",
) -> int:
    if source is None:
        source = file.filename or f"uploaded-{uuid4()}"

    text = await extract_text_from_upload(file)
    if not text.strip():
        raise ValueError("Uploaded document contained no text to ingest.")

    chunks = chunk_text(text)
    if not chunks:
        raise ValueError("Could not split uploaded document into chunks for ingestion.")

    ids = [
        _make_chunk_id(source, idx, chunk)
        for idx, chunk in enumerate(chunks)
    ]

    metadata = metadata or {}
    
    documents = [
        Document(
            page_content=chunk,
            metadata={
                **metadata,
                "source": source,
                "chunk_index": idx,
            },
        )
        for idx, chunk in enumerate(chunks)
    ]

    vector_store = get_vector_store(collection_name)

    vector_store.add_documents(
        documents=documents,
        ids=ids,
    )

    return len(chunks)
